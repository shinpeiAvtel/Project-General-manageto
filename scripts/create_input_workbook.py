from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Input" / "Personal_Schedule.xlsx"

HEADERS = [
    "Date", "Project", "Site", "Schedule / Task", "Start Time", "End Time",
    "Status", "Remarks", "Category", "Priority", "All Day",
    "Remote / Site / Office", "Last Updated"
]

PEOPLE = {
    "Shiraishi": [
        [date(2026, 9, 14), "Project A", "Tokyo", "Site Meeting", "09:00", "12:00", "Confirmed", "Kickoff meeting", "Meeting", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 15), "Project B", "Chiba", "Installation", "09:00", "17:00", "Confirmed", "CCTV installation", "Installation", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 16), "Project A", "Tokyo", "T&C", "10:00", "16:00", "Planned", "Client attendance", "T&C", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 16), "Project C", "Tokyo", "Design Review", "15:00", "17:00", "Planned", "Intentional overlap sample", "Meeting", "Medium", False, "Office", date(2026, 9, 10)],
        [date(2026, 9, 18), "Project C", "Tokyo", "Documentation", "09:00", "15:00", "Planned", "As-built review", "Documentation", "Medium", False, "Office", date(2026, 9, 10)],
    ],
    "Tanaka": [
        [date(2026, 9, 14), "Project C", "Tokyo", "Installation", "09:00", "17:00", "Confirmed", "", "Installation", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 15), "Project C", "Tokyo", "Installation", "09:00", "17:00", "Confirmed", "", "Installation", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 16), "Project D", "Yokohama", "Testing", "10:00", "16:00", "Planned", "", "Testing", "Medium", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 17), "Project D", "Yokohama", "Testing", "10:00", "16:00", "Planned", "", "Testing", "Medium", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 18), "Internal", "Tokyo", "Weekly Report", "15:00", "17:00", "Planned", "", "Internal", "Low", False, "Office", date(2026, 9, 10)],
    ],
    "Suzuki": [
        [date(2026, 9, 14), "Internal", "Tokyo", "Office Work", "09:00", "18:00", "Confirmed", "", "Internal", "Low", False, "Office", date(2026, 9, 10)],
        [date(2026, 9, 15), "Project A", "Tokyo", "Programming", "09:00", "17:00", "Confirmed", "", "Programming", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 16), "Project A", "Tokyo", "Programming", "09:00", "17:00", "Confirmed", "", "Programming", "High", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 17), "Project B", "Chiba", "Inspection", "10:00", "15:00", "Planned", "", "Inspection", "Medium", False, "Site", date(2026, 9, 10)],
        [date(2026, 9, 18), "LEAVE", "N/A", "Annual Leave", "", "", "Leave", "", "Leave", "Low", True, "Remote", date(2026, 9, 10)],
    ],
}

SYSTEM_SHEETS = ["00_Master_Schedule", "01_Monthly_View", "02_Weekly_View", "03_Conflict_Report", "04_Availability", "Template"]


def format_sheet(ws):
    for i, header in enumerate(HEADERS, 1):
        c = ws.cell(1, i, header)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:M1"
    widths = [13, 20, 16, 30, 11, 11, 13, 30, 16, 11, 11, 20, 14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    status_dv = DataValidation(type="list", formula1='"Planned,Confirmed,Completed,Cancelled,Leave"', allow_blank=True)
    priority_dv = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=True)
    mode_dv = DataValidation(type="list", formula1='"Remote,Site,Office"', allow_blank=True)
    ws.add_data_validation(status_dv); status_dv.add("G2:G500")
    ws.add_data_validation(priority_dv); priority_dv.add("J2:J500")
    ws.add_data_validation(mode_dv); mode_dv.add("L2:L500")
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 6), min_col=1, max_col=13):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r in range(2, max(ws.max_row, 6) + 1):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 13).number_format = "yyyy-mm-dd"


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = wb.active
    first.title = SYSTEM_SHEETS[0]
    for name in SYSTEM_SHEETS[1:]:
        wb.create_sheet(name)
    for name in SYSTEM_SHEETS[:-1]:
        ws = wb[name]
        ws["A1"] = "AUTO GENERATED - DO NOT EDIT MANUALLY"
        ws["A1"].font = Font(bold=True)
        ws["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    format_sheet(wb["Template"])
    for person, rows in PEOPLE.items():
        ws = wb.create_sheet(person)
        ws.append(HEADERS)
        for row in rows:
            ws.append(row)
        format_sheet(ws)
    wb.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Persons: {len(PEOPLE)}")
    print(f"Records: {sum(len(v) for v in PEOPLE.values())}")


if __name__ == "__main__":
    main()
