from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from .build_monthly_view import build_monthly_view_sheet
from .build_weekly_view import build_weekly_view_sheet
from .calculate_availability import calculate_availability
from .detect_conflicts import detect_conflicts
from .utils import (
    MASTER_COLUMNS,
    STATUS_FILLS,
    apply_date_time_formats,
    apply_status_formatting,
    apply_table_style,
    load_schedule_config,
    record_sort_key,
    repo_root,
    replace_sheet,
    save_json,
    style_header_row,
    write_banner,
)
from .validate_schedule import run_validation


CONFLICT_COLUMNS = [
    "Person",
    "Date",
    "Schedule A",
    "Schedule B",
    "Project A",
    "Project B",
    "Time A",
    "Time B",
    "Conflict Type",
    "Source Sheet A",
    "Source Row A",
    "Source Sheet B",
    "Source Row B",
]

AVAILABILITY_COLUMNS = [
    "Date",
    "Person",
    "Availability",
    "Scheduled Minutes",
    "Working Minutes",
    "Notes",
]


def build_master_sheet(workbook: Workbook, records: list[dict[str, Any]]):
    worksheet = replace_sheet(workbook, "00_Master_Schedule", 0)
    write_banner(worksheet, len(MASTER_COLUMNS))
    for column_index, header in enumerate(MASTER_COLUMNS, start=1):
        worksheet.cell(row=2, column=column_index, value=header)

    for row_index, record in enumerate(sorted(records, key=record_sort_key), start=3):
        ordered_values = [
            record.get("date"),
            record.get("day"),
            record.get("person"),
            record.get("project"),
            record.get("site"),
            record.get("schedule_task"),
            record.get("start_time"),
            record.get("end_time"),
            record.get("status"),
            record.get("category"),
            record.get("priority"),
            record.get("remarks"),
            "TRUE" if record.get("all_day") else "FALSE",
            record.get("source_sheet"),
            record.get("source_row"),
            record.get("last_updated"),
        ]
        for column_index, value in enumerate(ordered_values, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)

    style_header_row(worksheet, 2)
    apply_table_style(worksheet, header_row=2, freeze_panes="A3")
    apply_date_time_formats(worksheet, header_row=2)
    apply_status_formatting(worksheet, 2, "Status", {key: value for key, value in STATUS_FILLS.items() if key in {"Planned", "Confirmed", "Completed", "Cancelled", "Leave"}})
    return worksheet


def build_conflict_sheet(workbook: Workbook, conflicts: list[dict[str, Any]]):
    worksheet = replace_sheet(workbook, "03_Conflict_Report", 3)
    write_banner(worksheet, len(CONFLICT_COLUMNS))
    for column_index, header in enumerate(CONFLICT_COLUMNS, start=1):
        worksheet.cell(row=2, column=column_index, value=header)
    for row_index, conflict in enumerate(conflicts, start=3):
        values = [
            conflict.get("person"),
            conflict.get("date"),
            conflict.get("schedule_a"),
            conflict.get("schedule_b"),
            conflict.get("project_a"),
            conflict.get("project_b"),
            conflict.get("time_a"),
            conflict.get("time_b"),
            conflict.get("conflict_type"),
            conflict.get("source_sheet_a"),
            conflict.get("source_row_a"),
            conflict.get("source_sheet_b"),
            conflict.get("source_row_b"),
        ]
        for column_index, value in enumerate(values, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)
    style_header_row(worksheet, 2)
    apply_table_style(worksheet, header_row=2, freeze_panes="A3")
    apply_date_time_formats(worksheet, header_row=2)
    apply_status_formatting(worksheet, 2, "Conflict Type", {"CONFLICT": STATUS_FILLS["CONFLICT"]})
    return worksheet


def build_availability_sheet(workbook: Workbook, availability_rows: list[dict[str, Any]]):
    worksheet = replace_sheet(workbook, "04_Availability", 4)
    write_banner(worksheet, len(AVAILABILITY_COLUMNS))
    for column_index, header in enumerate(AVAILABILITY_COLUMNS, start=1):
        worksheet.cell(row=2, column=column_index, value=header)
    for row_index, availability in enumerate(availability_rows, start=3):
        values = [
            availability.get("date"),
            availability.get("person"),
            availability.get("availability"),
            availability.get("scheduled_minutes"),
            availability.get("working_minutes"),
            availability.get("notes"),
        ]
        for column_index, value in enumerate(values, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)
    style_header_row(worksheet, 2)
    apply_table_style(worksheet, header_row=2, freeze_panes="A3")
    apply_date_time_formats(worksheet, header_row=2)
    apply_status_formatting(
        worksheet,
        2,
        "Availability",
        {key: STATUS_FILLS[key] for key in ["AVAILABLE", "PARTIAL", "BOOKED", "LEAVE"]},
    )
    return worksheet


def build_output_artifacts(config_path: str | Path | None = None) -> dict[str, Any]:
    config = load_schedule_config(config_path)
    validation_result = run_validation(config_path)
    records = validation_result["records"]
    issues = validation_result["issues"]
    conflicts = detect_conflicts(records)
    availability_rows = calculate_availability(records, config)

    base_date_value = config["query"].get("base_date")
    base_date = date.fromisoformat(base_date_value) if base_date_value else date.today()

    workbook = Workbook()
    workbook.remove(workbook.active)
    build_master_sheet(workbook, records)
    build_monthly_view_sheet(workbook, records)
    build_weekly_view_sheet(workbook, records, base_date=base_date)
    build_conflict_sheet(workbook, conflicts)
    build_availability_sheet(workbook, availability_rows)

    output_path = repo_root() / config["workbook"]["output_path"]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    save_json(
        repo_root() / "data" / "build_summary.json",
        {
            "generated_at": date.today().isoformat(),
            "input_workbook": config["workbook"]["input_path"],
            "output_workbook": config["workbook"]["output_path"],
            "detected_personal_sheets": validation_result["detected_personal_sheets"],
            "record_count": len(records),
            "issue_count": len(issues),
            "conflict_count": len(conflicts),
            "availability_rows": len(availability_rows),
            "base_date": base_date.isoformat(),
        },
    )

    return {
        "output_path": str(output_path),
        "records": records,
        "issues": issues,
        "conflicts": conflicts,
        "availability_rows": availability_rows,
        "detected_personal_sheets": validation_result["detected_personal_sheets"],
        "base_date": base_date.isoformat(),
    }


def main() -> None:
    result = build_output_artifacts()
    print(f"Generated workbook: {result['output_path']}")
    print(f"Personal sheets: {', '.join(result['detected_personal_sheets'])}")
    print(f"Records: {len(result['records'])} | Conflicts: {len(result['conflicts'])} | Availability rows: {len(result['availability_rows'])}")


if __name__ == "__main__":
    main()
