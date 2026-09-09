from __future__ import annotations

from pathlib import Path
from typing import Any
from openpyxl import Workbook

from scripts.utils import ensure_dir


def generate_master_workbook(
    normalized_rows: list[dict[str, Any]],
    analysis: dict[str, Any],
    validation: dict[str, Any],
    output_dir: Path,
) -> Path:
    ensure_dir(output_dir)
    wb = Workbook()

    ws_master = wb.active
    ws_master.title = "Master Schedule"
    master_headers = [
        "Project Name",
        "Task ID",
        "Task Name",
        "Start Date",
        "End Date",
        "Progress",
        "Status",
        "Delay Status",
        "Days Remaining",
        "Assigned Person",
        "Project Manager",
        "Milestone",
        "Priority",
        "Source File",
        "Source Sheet",
    ]
    ws_master.append(master_headers)
    for row in normalized_rows:
        ws_master.append(
            [
                row.get("project_name"),
                row.get("task_id"),
                row.get("task_name"),
                row.get("start_date"),
                row.get("end_date"),
                row.get("progress"),
                row.get("status"),
                row.get("delay_status"),
                row.get("days_remaining"),
                ", ".join(row.get("assigned_person", [])),
                row.get("project_manager"),
                row.get("milestone"),
                row.get("priority"),
                row.get("source_file"),
                row.get("source_sheet"),
            ]
        )

    ws_summary = wb.create_sheet("Project Summary")
    summary_headers = [
        "Project Name",
        "Project Manager",
        "Start Date",
        "End Date",
        "Total Tasks",
        "Completed Tasks",
        "Progress",
        "Delayed Tasks",
        "Milestones",
        "Last Updated",
        "Risk Level",
        "Risk Score",
    ]
    ws_summary.append(summary_headers)
    for item in analysis.get("project_summary", []):
        ws_summary.append(
            [
                item.get("project_name"),
                item.get("project_manager"),
                item.get("start_date"),
                item.get("end_date"),
                item.get("total_tasks"),
                item.get("completed_tasks"),
                item.get("progress"),
                item.get("delayed_tasks"),
                item.get("milestones"),
                item.get("last_updated_days_ago"),
                item.get("risk_level"),
                item.get("risk_score"),
            ]
        )

    ws_milestones = wb.create_sheet("Milestones")
    ws_milestones.append(["Project Name", "Task ID", "Task Name", "End Date", "Category"])
    for row in analysis.get("milestones", []):
        ws_milestones.append([row.get("project_name"), row.get("task_id"), row.get("task_name"), row.get("end_date"), row.get("category")])

    ws_delayed = wb.create_sheet("Delayed Tasks")
    ws_delayed.append(["Project Name", "Task ID", "Task Name", "End Date", "Progress", "Delay Status"])
    for row in analysis.get("delayed_tasks", []):
        ws_delayed.append([row.get("project_name"), row.get("task_id"), row.get("task_name"), row.get("end_date"), row.get("progress"), row.get("delay_status")])

    ws_conflicts = wb.create_sheet("Resource Conflicts")
    ws_conflicts.append(["Person", "Date", "Projects", "Status"])
    for c in analysis.get("resource_conflicts", []):
        ws_conflicts.append([c.get("person"), c.get("date"), ", ".join(c.get("projects", [])), c.get("status")])

    ws_upcoming = wb.create_sheet("Upcoming 14 Days")
    ws_upcoming.append(["Project Name", "Task ID", "Task Name", "Start Date", "End Date", "Priority", "Milestone", "Delay Status", "Category"])
    for r in analysis.get("upcoming_14_days", []):
        ws_upcoming.append(
            [
                r.get("project_name"),
                r.get("task_id"),
                r.get("task_name"),
                r.get("start_date"),
                r.get("end_date"),
                r.get("priority"),
                r.get("milestone"),
                r.get("delay_status"),
                r.get("category"),
            ]
        )

    ws_ve = wb.create_sheet("Validation Errors")
    ws_ve.append(["Severity", "Code", "Field", "Message", "Project Name", "Task ID", "Source File", "Source Sheet"])
    for i in validation.get("issues", []):
        ws_ve.append([i.get("severity"), i.get("code"), i.get("field"), i.get("message"), i.get("project_name"), i.get("task_id"), i.get("source_file"), i.get("source_sheet")])

    out = output_dir / "Master_Schedule.xlsx"
    wb.save(out)
    return out
