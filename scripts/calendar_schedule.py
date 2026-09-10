from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
import calendar
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

INPUT = Path('Input/Personal_Schedule.xlsx')
OUTPUT = Path('Output/Master_Schedule.xlsx')
SYSTEM_SHEETS = {
    '00_Master_Schedule', '01_Monthly_View', '02_Weekly_View',
    '03_Conflict_Report', '04_Availability', '05_Validation_Report',
    'Calendar_Template', 'Config'
}
DAY_NAMES = ['Mon','Tue','Wed','Thu','Fri','Sat','Sun']

HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
WEEKEND_FILL = PatternFill('solid', fgColor='FCE4D6')
WEEKDAY_FILL = PatternFill('solid', fgColor='D9EAF7')
OUT_MONTH_FILL = PatternFill('solid', fgColor='F2F2F2')
WHITE_FONT = Font(color='FFFFFF', bold=True)
THIN = Side(style='thin', color='D9D9D9')


def _month_start(year: int, month: int) -> date:
    first = date(year, month, 1)
    return first - timedelta(days=first.weekday())


def build_calendar_sheet(ws, person: str, year: int, month: int) -> None:
    ws.merge_cells('A1:U1')
    ws['A1'] = f'{person} Schedule Calendar'
    ws['A1'].fill = HEADER_FILL
    ws['A1'].font = Font(color='FFFFFF', bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws['A2'], ws['B2'] = 'Year', year
    ws['A3'], ws['B3'] = 'Month', month
    ws.merge_cells('D2:U3')
    ws['D2'] = '各日の Time / Project / Schedule 欄に入力してください。1日3件まで。時間は HH:MM-HH:MM、終日は ALL DAY。'
    ws['D2'].alignment = Alignment(wrap_text=True, vertical='center')
    ws['D2'].fill = PatternFill('solid', fgColor='EAF2F8')

    for d in range(7):
        ws.column_dimensions[get_column_letter(d*3+1)].width = 12
        ws.column_dimensions[get_column_letter(d*3+2)].width = 16
        ws.column_dimensions[get_column_letter(d*3+3)].width = 24

    time_dv = DataValidation(type='list', formula1='"ALL DAY,08:00-12:00,09:00-12:00,09:00-17:00,09:00-18:00,10:00-15:00,10:00-16:00,13:00-17:00,15:00-17:00"')
    ws.add_data_validation(time_dv)

    start = _month_start(year, month)
    row = 5
    for week in range(6):
        for d in range(7):
            current = start + timedelta(days=week*7+d)
            c = d*3+1
            ws.merge_cells(start_row=row, start_column=c, end_row=row, end_column=c+2)
            head = ws.cell(row=row, column=c)
            head.value = f'{DAY_NAMES[d]}  {current.day}'
            head.font = Font(bold=True)
            head.alignment = Alignment(horizontal='center')
            if current.month != month:
                head.fill = OUT_MONTH_FILL
            elif d >= 5:
                head.fill = WEEKEND_FILL
            else:
                head.fill = WEEKDAY_FILL
            for cc in range(c, c+3):
                ws.cell(row=row, column=cc).border = Border(top=THIN,bottom=THIN,left=THIN,right=THIN)
            for slot in range(3):
                rr = row + 1 + slot
                for cc in range(c, c+3):
                    cell = ws.cell(rr, cc)
                    cell.border = Border(top=THIN,bottom=THIN,left=THIN,right=THIN)
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                time_dv.add(ws.cell(rr, c))
        row += 5
    ws.freeze_panes = 'A5'


def create_input(year: int = 2026, month: int = 9) -> None:
    if INPUT.exists():
        return
    INPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name in ['Calendar_Template','Shiraishi','Tanaka','Suzuki']:
        ws = wb.create_sheet(name)
        build_calendar_sheet(ws, name, year, month)
    wb.save(INPUT)


def parse_calendar_sheet(ws):
    year = int(ws['B2'].value)
    month = int(ws['B3'].value)
    start = _month_start(year, month)
    records = []
    row = 5
    for week in range(6):
        for d in range(7):
            current = start + timedelta(days=week*7+d)
            c = d*3+1
            for slot in range(3):
                rr = row + 1 + slot
                time_text = str(ws.cell(rr,c).value or '').strip()
                project = str(ws.cell(rr,c+1).value or '').strip()
                task = str(ws.cell(rr,c+2).value or '').strip()
                if not any([time_text, project, task]):
                    continue
                records.append({
                    'date': current,
                    'person': ws.title,
                    'time': time_text,
                    'project': project or 'UNKNOWN',
                    'task': task or 'UNKNOWN',
                    'source_sheet': ws.title,
                    'source_row': rr,
                })
        row += 5
    return records


def _minutes(t: str):
    if not t or t == 'ALL DAY' or '-' not in t:
        return None
    m = re.fullmatch(r'(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})', t)
    if not m:
        return None
    sh, sm, eh, em = map(int, m.groups())
    return sh*60+sm, eh*60+em


def detect_conflicts(records):
    out=[]
    by_key={}
    for r in records:
        by_key.setdefault((r['person'], r['date']), []).append(r)
    for (person, dt), items in by_key.items():
        for i,a in enumerate(items):
            ra=_minutes(a['time'])
            if not ra: continue
            for b in items[i+1:]:
                rb=_minutes(b['time'])
                if not rb: continue
                if max(ra[0],rb[0]) < min(ra[1],rb[1]):
                    out.append((person,dt,a,b))
    return out


def build_master() -> None:
    create_input()
    src = load_workbook(INPUT, data_only=False)
    records=[]
    for ws in src.worksheets:
        if ws.title in SYSTEM_SHEETS:
            continue
        records.extend(parse_calendar_sheet(ws))
    records.sort(key=lambda r:(r['date'], r['time'], r['person']))

    wb=Workbook()
    wb.remove(wb.active)
    master=wb.create_sheet('00_Master_Schedule')
    headers=['Date','Day','Person','Project','Schedule / Task','Time','Source Sheet','Source Row']
    master.append(headers)
    for c in master[1]:
        c.fill=HEADER_FILL; c.font=WHITE_FONT
    for r in records:
        master.append([r['date'],r['date'].strftime('%a'),r['person'],r['project'],r['task'],r['time'],r['source_sheet'],r['source_row']])
    master.freeze_panes='A2'
    master.auto_filter.ref=master.dimensions
    for cell in master['A'][1:]: cell.number_format='yyyy-mm-dd'

    monthly=wb.create_sheet('01_Monthly_View')
    dates=sorted({r['date'] for r in records})
    people=sorted({r['person'] for r in records})
    monthly.append(['Person']+[d.isoformat() for d in dates])
    for c in monthly[1]: c.fill=HEADER_FILL; c.font=WHITE_FONT
    for p in people:
        row=[p]
        for d in dates:
            items=[r for r in records if r['person']==p and r['date']==d]
            row.append('\n'.join(f"{x['time']} | {x['project']} | {x['task']}" for x in items))
        monthly.append(row)
    for row in monthly.iter_rows():
        for cell in row: cell.alignment=Alignment(wrap_text=True, vertical='top')

    weekly=wb.create_sheet('02_Weekly_View')
    weekly.append(headers)
    for c in weekly[1]: c.fill=HEADER_FILL; c.font=WHITE_FONT
    for r in records:
        weekly.append([r['date'],r['date'].strftime('%a'),r['person'],r['project'],r['task'],r['time'],r['source_sheet'],r['source_row']])

    conf=wb.create_sheet('03_Conflict_Report')
    conf_headers=['Person','Date','Project A','Schedule A','Time A','Project B','Schedule B','Time B','Conflict Type']
    conf.append(conf_headers)
    for c in conf[1]: c.fill=HEADER_FILL; c.font=WHITE_FONT
    for person,dt,a,b in detect_conflicts(records):
        conf.append([person,dt,a['project'],a['task'],a['time'],b['project'],b['task'],b['time'],'TIME_OVERLAP'])

    avail=wb.create_sheet('04_Availability')
    avail.append(['Person']+[d.isoformat() for d in dates])
    for c in avail[1]: c.fill=HEADER_FILL; c.font=WHITE_FONT
    for p in people:
        row=[p]
        for d in dates:
            items=[r for r in records if r['person']==p and r['date']==d]
            if not items: state='AVAILABLE'
            elif any(r['time']=='ALL DAY' for r in items): state='BOOKED'
            else: state='PARTIAL'
            row.append(state)
        avail.append(row)

    val=wb.create_sheet('05_Validation_Report')
    val.append(['Severity','Person','Date','Source Row','Field','Message'])
    for c in val[1]: c.fill=HEADER_FILL; c.font=WHITE_FONT
    for r in records:
        if r['time'] and r['time']!='ALL DAY' and not _minutes(r['time']):
            val.append(['ERROR',r['person'],r['date'],r['source_row'],'Time','Use HH:MM-HH:MM or ALL DAY'])
        if r['project']=='UNKNOWN':
            val.append(['WARNING',r['person'],r['date'],r['source_row'],'Project','Project is missing'])
        if r['task']=='UNKNOWN':
            val.append(['WARNING',r['person'],r['date'],r['source_row'],'Schedule / Task','Schedule is missing'])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f'persons={len(people)} records={len(records)} conflicts={len(detect_conflicts(records))} output={OUTPUT}')


if __name__ == '__main__':
    build_master()
