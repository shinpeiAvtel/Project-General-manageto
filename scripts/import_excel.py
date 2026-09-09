from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from scripts.utils import ensure_dir, load_yaml
except ModuleNotFoundError:  # pragma: no cover
    from utils import ensure_dir, load_yaml


class ExcelImporter:
    def __init__(self, schema_path: str | Path):
        schema = load_yaml(schema_path)
        self.required_headers = schema.get("required_headers", [])

    def import_file(self, file_path: str | Path, input_root: str | Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        file_path = Path(file_path)
        rel = file_path.relative_to(Path(input_root))
        parts = rel.parts
        folder_owner = parts[0] if len(parts) >= 3 else "MISSING"
        folder_project = parts[1] if len(parts) >= 3 else "MISSING"

        wb = load_workbook(file_path, read_only=True, data_only=True)
        imported_at = datetime.now(timezone.utc).isoformat()
        all_records: list[dict[str, Any]] = []
        issues: list[dict[str, Any]] = []

        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [str(c).strip() if c is not None else "" for c in header_row]
            missing_headers = [h for h in self.required_headers if h not in headers]
            if missing_headers:
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "MISSING_HEADER",
                        "message": f"Missing headers: {', '.join(missing_headers)}",
                        "source_file": str(file_path),
                        "source_sheet": ws.title,
                    }
                )
                continue

            for idx, row in enumerate(rows, start=2):
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                record = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}
                record.update(
                    {
                        "_source_file": str(file_path),
                        "_source_sheet": ws.title,
                        "_source_row": idx,
                        "_imported_at": imported_at,
                        "_folder_owner": folder_owner,
                        "_folder_project": folder_project,
                    }
                )
                all_records.append(record)

        return all_records, issues


def discover_xlsx_files(input_dir: str | Path) -> list[Path]:
    return sorted([p for p in Path(input_dir).rglob("*.xlsx") if not p.name.startswith("~$")])


def main() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    schema_path = repo_root / "config" / "schedule_schema.yaml"
    input_root = repo_root / "Input"
    importer = ExcelImporter(schema_path)

    records: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for xlsx in discover_xlsx_files(input_root):
        recs, errs = importer.import_file(xlsx, input_root)
        records.extend(recs)
        issues.extend(errs)

    out_dir = repo_root / "data" / "raw"
    ensure_dir(out_dir)
    (out_dir / "imported_records.json").write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (out_dir / "import_issues.json").write_text(json.dumps(issues, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Imported {len(records)} records with {len(issues)} import issues")


if __name__ == "__main__":
    main()
