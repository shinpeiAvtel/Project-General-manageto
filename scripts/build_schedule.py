from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = ROOT / "Input" / "Personal_Schedule.xlsx"
CONFIG_FILE = ROOT / "config" / "schedule_config.yaml"
OUTPUT_FILE = ROOT / "Output" / "Master_Schedule.xlsx"

HEADERS = [
    "Date", "Project", "Site", "Schedule / Task", "Start Time", "End Time",
    "Status", "Remarks", "Category", "Priority", "All Day",
    "Remote / Site / Office", "Last Updated"
]
MASTER_HEADERS = [
    "Date", "Day", "Person", "Project", "Site", "Schedule / Task",
    "Start Time", "End Time", "Status", "Category", "Priority", "Remarks",
    "Source Sheet", "Source Row", "Last Updated"
]
VALIDATION_HEADERS = ["Severity", "Person", "Sheet", "Row", "Field", "Message"]


def load_config() -> dict[str, Any]:
    with CONFIG_FILE.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def normalize_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def normalize_time(value: Any) -> time | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, str):
        for fmt in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(value.strip(), fmt).time()
            except ValueError:
                pass
    return None


def to_minutes(t: time) -> int:
    return t.hour * 60 + t.minute


def create_template() -> Path:
    INPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "00_Master_Schedule"
    for name in ["01_Monthly_View", "02_Weekly_View", "03_Conflict_Report", "04_Availability", "Template"]:
        wb.create_sheet(name)

    people = {
        "Shiraishi": [
            [date(2026, 9, 14), "Project A", "Tokyo", "Site Meeting", "09:00", "12:00", "Confirmed", "Kickoff meeting", "Meeting", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 15), "Project B", "Chiba", "Installation", "09:00", "17:00", "Confirmed", "CCTV installation", "Installation", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 16), "Project A", "Tokyo", "T&C", "10:00", "16:00", "Planned", "Client attendance", "T&C", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 16), "Project C", "Tokyo", "Design Review", "15:00", "17:00", "Planned", "Intentional overlap sample", "Meeting", "Medium", False, "Office", date(2026, 9, 10)],
            [date(2026, 9, 18), "Project C", "Tokyo", "Documentation", "09:00", "15:00", "Planned", "As-built review", "Documentation", "Medium", False, "Office", date(2026, 9, 10)],
        ],
        "Tanaka": [
            [date(2026, 9, 14), "Project C", "Tokyo", "Installation", "09:00", "17:00", "Confirmed", "", "Installation", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 15), "Project C", "Tokyo", "Installation", "09:00", "17:00", "Confirmed", "", "Installation", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 16), "Project D", "Yokohama", "Testing", "10:00", "16:00", "Planned", "", "Testing", "Medium", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 17), "Project D", "Yokohama", "Testing", "10:00", "16:00", "Planned", "", "Testing", "Medium", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 18), "Internal", "Tokyo", "Weekly Report", "15:00", "17:00", "Planned", "", "Internal", "Low", False, "Office", date(2026, 9, 10)],
        ],
        "Suzuki": [
            [date(2026, 9, 14), "Internal", "Tokyo", "Office Work", "09:00", "18:00", "Confirmed", "", "Internal", "Low", False, "Office", date(2026, 9, 10)],
            [date(2026, 9, 15), "Project A", "Tokyo", "Programming", "09:00", "17:00", "Confirmed", "", "Programming", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 16), "Project A", "Tokyo", "Programming", "09:00", "17:00", "Confirmed", "", "Programming", "High", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 17), "Project B", "Chiba", "Inspection", "10:00", "15:00", "Planned", "", "Inspection", "Medium", False, "Site", date(2026, 9, 10)],
            [date(2026, 9, 18), "LEAVE", "N/A", "Annual Leave", "", "", "Leave", "", "Leave", "Low", True, "Remote", date(2026, 9, 10)],
        ],
    }

    style_header(wb["Template"], HEADERS)
    setup_person_sheet(wb["Template"])
    for person, rows in people.items():
        ps = wb.create_sheet(person)
        style_header(ps, HEADERS)
        setup_person_sheet(ps)
        for r in rows:
            ps.append(r)

    for name in ["00_Master_Schedule", "01_Monthly_View", "02_Weekly_View", "03_Conflict_Report", "04_Availability"]:
        wb[name]["A1"] = "AUTO GENERATED - DO NOT EDIT MANUALLY"
        wb[name]["A1"].font = Font(bold=True)
        wb[name]["A1"].fill = PatternFill("solid", fgColor="D9EAF7")

    wb.save(INPUT_FILE)
    return INPUT_FILE


def style_header(ws, headers: list[str]) -> None:
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(1, idx, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def setup_person_sheet(ws) -> None:
    widths = [13, 20, 16, 30, 11, 11, 13, 30, 16, 11, 11, 20, 14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    status_dv = DataValidation(type="list", formula1='"Planned,Confirmed,Completed,Cancelled,Leave"', allow_blank=True)
    priority_dv = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=True)
    mode_dv = DataValidation(type="list", formula1='"Remote,Site,Office"', allow_blank=True)
    ws.add_data_validation(status_dv); status_dv.add("G2:G500")
    ws.add_data_validation(priority_dv); priority_dv.add("J2:J500")
    ws.add_data_validation(mode_dv); mode_dv.add("L2:L500")
    for row in range(2, 501):
        ws.cell(row, 1).number_format = "yyyy-mm-dd"
        ws.cell(row, 13).number_format = "yyyy-mm-dd"
        for col in range(1, 14):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)


def read_records(wb, config: dict[str, Any]) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    system_sheets = set(config["system_sheets"])
    statuses = set(config["statuses"])
    records: list[dict[str, Any]] = []
    errors: list[list[Any]] = []

    for ws in wb.worksheets:
        if ws.title in system_sheets:
            continue
        person = ws.title.strip()
        actual_headers = [ws.cell(1, c).value for c in range(1, len(HEADERS) + 1)]
        if actual_headers != HEADERS:
            errors.append(["ERROR", person, ws.title, 1, "Header", "Header does not match the required template"])
            continue

        for row in range(2, ws.max_row + 1):
            values = [ws.cell(row, c).value for c in range(1, len(HEADERS) + 1)]
            if all(v in (None, "") for v in values):
                continue
            raw = dict(zip(HEADERS, values))
            d = normalize_date(raw["Date"])
            st = normalize_time(raw["Start Time"])
            et = normalize_time(raw["End Time"])
            all_day = bool(raw["All Day"])
            status = str(raw["Status"] or "").strip()

            if not d:
                errors.append(["ERROR", person, ws.title, row, "Date", "Missing or invalid Date (YYYY-MM-DD required)"])
            if not raw["Schedule / Task"]:
                errors.append(["ERROR", person, ws.title, row, "Schedule / Task", "Schedule / Task is required"])
            if status and status not in statuses:
                errors.append(["WARNING", person, ws.title, row, "Status", f"Unknown Status: {status}"])
            if not all_day and status != "Leave":
                if not st or not et:
                    errors.append(["ERROR", person, ws.title, row, "Time", "Start Time and End Time are required unless All Day/Leave"])
                elif to_minutes(st) >= to_minutes(et):
                    errors.append(["ERROR", person, ws.title, row, "Time", "Start Time must be earlier than End Time"])
            if not d or not raw["Schedule / Task"]:
                continue

            records.append({
                "date": d,
                "person": person,
                "project": raw["Project"] or "UNKNOWN",
                "site": raw["Site"] or "UNKNOWN",
                "task": raw["Schedule / Task"],
                "start": st,
                "end": et,
                "status": status or "UNKNOWN",
                "remarks": raw["Remarks"] or "",
                "category": raw["Category"] or "UNKNOWN",
                "priority": raw["Priority"] or "UNKNOWN",
                "all_day": all_day,
                "mode": raw["Remote / Site / Office"] or "UNKNOWN",
                "last_updated": normalize_date(raw["Last Updated"]),
                "source_sheet": ws.title,
                "source_row": row,
            })
    return records, errors


def detect_conflicts(records: list[dict[str, Any]]) -> list[list[Any]]:
    grouped = defaultdict(list)
    for r in records:
        if r["status"] == "Cancelled":
            continue
        grouped[(r["person"], r["date"])].append(r)
    out = []
    for (person, d), items in grouped.items():
        for i in range(len(items)):
            for j in range(i + 1, len(items)):
                a, b = items[i], items[j]
                conflict = False
                if a["all_day"] or b["all_day"] or a["status"] == "Leave" or b["status"] == "Leave":
                    conflict = True
                elif a["start"] and a["end"] and b["start"] and b["end"]:
                    conflict = max(to_minutes(a["start"]), to_minutes(b["start"])) < min(to_minutes(a["end"]), to_minutes(b["end"]))
                if conflict:
                    out.append([
                        person, d, a["project"], a["task"], fmt_range(a),
                        b["project"], b["task"], fmt_range(b), "TIME_OVERLAP"
                    ])
    return out


def fmt_time(t: time | None) -> str:
    return t.strftime("%H:%M") if t else ""


def fmt_range(r: dict[str, Any]) -> str:
    if r["all_day"] or r["status"] == "Leave":
        return "ALL DAY"
    return f"{fmt_time(r['start'])}-{fmt_time(r['end'])}"


def build_output(records: list[dict[str, Any]], errors: list[list[Any]], config: dict[str, Any]) -> Path:
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    master = wb.active
    master.title = "00_Master_Schedule"
    monthly = wb.create_sheet("01_Monthly_View")
    weekly = wb.create_sheet("02_Weekly_View")
    conflicts_ws = wb.create_sheet("03_Conflict_Report")
    availability_ws = wb.create_sheet("04_Availability")
    validation_ws = wb.create_sheet("05_Validation_Report")

    records = sorted(records, key=lambda r: (r["date"], to_minutes(r["start"]) if r["start"] else -1, r["person"]))
    style_header(master, MASTER_HEADERS)
    for r in records:
        master.append([
            r["date"], r["date"].strftime("%a"), r["person"], r["project"], r["site"], r["task"],
            fmt_time(r["start"]), fmt_time(r["end"]), r["status"], r["category"], r["priority"], r["remarks"],
            r["source_sheet"], r["source_row"], r["last_updated"]
        ])
    style_output(master, [13, 8, 16, 20, 16, 30, 11, 11, 13, 16, 11, 30, 16, 11, 14])

    build_monthly(monthly, records)
    build_weekly(weekly, records)
    conflicts = detect_conflicts(records)
    conflict_headers = ["Person", "Date", "Project A", "Schedule A", "Time A", "Project B", "Schedule B", "Time B", "Conflict Type"]
    style_header(conflicts_ws, conflict_headers)
    for row in conflicts:
        conflicts_ws.append(row)
    style_output(conflicts_ws, [16, 13, 18, 26, 15, 18, 26, 15, 18])

    build_availability(availability_ws, records, config)
    style_header(validation_ws, VALIDATION_HEADERS)
    for row in errors:
        validation_ws.append(row)
    style_output(validation_ws, [11, 16, 16, 9, 18, 50])

    wb.save(OUTPUT_FILE)
    return OUTPUT_FILE


def style_output(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for c in range(1, ws.max_column + 1):
        ws.cell(1, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.freeze_panes = "A2"


def build_monthly(ws, records: list[dict[str, Any]]) -> None:
    if not records:
        ws["A1"] = "No schedule data"
        return
    dates = sorted({r["date"] for r in records})
    people = sorted({r["person"] for r in records})
    ws.cell(1, 1, "Person")
    for c, d in enumerate(dates, 2):
        ws.cell(1, c, d)
        ws.cell(1, c).number_format = "m/d (ddd)"
    for r_idx, person in enumerate(people, 2):
        ws.cell(r_idx, 1, person)
        for c_idx, d in enumerate(dates, 2):
            items = [r for r in records if r["person"] == person and r["date"] == d and r["status"] != "Cancelled"]
            text = "\n".join(f"{x['project']} | {x['task']} | {fmt_range(x)}" for x in items)
            ws.cell(r_idx, c_idx, text)
    style_header(ws, ["Person"] + [d.strftime("%Y-%m-%d") for d in dates])
    ws.column_dimensions["A"].width = 18
    for c in range(2, len(dates) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 24
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_weekly(ws, records: list[dict[str, Any]]) -> None:
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    end = monday + timedelta(days=13)
    filtered = [r for r in records if monday <= r["date"] <= end]
    headers = ["Week", "Date", "Day", "Person", "Project", "Schedule / Task", "Time", "Status", "Site"]
    style_header(ws, headers)
    for r in filtered:
        ws.append([r["date"].isocalendar().week, r["date"], r["date"].strftime("%A"), r["person"], r["project"], r["task"], fmt_range(r), r["status"], r["site"]])
    style_output(ws, [9, 13, 12, 16, 20, 28, 16, 13, 16])


def build_availability(ws, records: list[dict[str, Any]], config: dict[str, Any]) -> None:
    people = sorted({r["person"] for r in records})
    dates = sorted({r["date"] for r in records})
    work_start = normalize_time(config["workday"]["start"])
    work_end = normalize_time(config["workday"]["end"])
    work_minutes = to_minutes(work_end) - to_minutes(work_start)
    ws.cell(1, 1, "Person")
    for c, d in enumerate(dates, 2):
        ws.cell(1, c, d)
        ws.cell(1, c).number_format = "m/d (ddd)"
    for r_idx, person in enumerate(people, 2):
        ws.cell(r_idx, 1, person)
        for c_idx, d in enumerate(dates, 2):
            items = [r for r in records if r["person"] == person and r["date"] == d and r["status"] != "Cancelled"]
            if any(r["status"] == "Leave" or r["all_day"] for r in items):
                state = "LEAVE" if any(r["status"] == "Leave" for r in items) else "BOOKED"
            else:
                booked = 0
                for r in items:
                    if r["start"] and r["end"]:
                        start = max(to_minutes(r["start"]), to_minutes(work_start))
                        end = min(to_minutes(r["end"]), to_minutes(work_end))
                        booked += max(0, end - start)
                if booked == 0:
                    state = "AVAILABLE"
                elif booked >= work_minutes:
                    state = "BOOKED"
                else:
                    state = "PARTIAL"
            ws.cell(r_idx, c_idx, state)
    style_header(ws, ["Person"] + [d.strftime("%Y-%m-%d") for d in dates])
    ws.column_dimensions["A"].width = 18
    for c in range(2, len(dates) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 15


def main() -> None:
    config = load_config()
    if not INPUT_FILE.exists():
        create_template()
        print(f"Created sample input: {INPUT_FILE}")
    wb = load_workbook(INPUT_FILE, data_only=False)
    records, errors = read_records(wb, config)
    output = build_output(records, errors, config)
    conflicts = detect_conflicts(records)
    people = sorted({r['person'] for r in records})
    print(f"Input: {INPUT_FILE}")
    print(f"Detected persons: {len(people)} -> {', '.join(people)}")
    print(f"Schedule records: {len(records)}")
    print(f"Conflicts: {len(conflicts)}")
    print(f"Validation messages: {len(errors)}")
    print(f"Output: {output}")


if __name__ == "__main__":
    main()
