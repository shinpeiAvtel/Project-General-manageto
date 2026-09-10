from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .utils import canonical_header, is_personal_sheet, load_schedule_config, normalise_text, repo_root

KNOWN_FIELDS = [
    "date",
    "project",
    "site",
    "schedule_task",
    "start_time",
    "end_time",
    "status",
    "remarks",
    "category",
    "client",
    "location",
    "priority",
    "all_day",
    "work_mode",
    "confirmed_state",
    "last_updated",
]

REQUIRED_FIELDS = {
    "date",
    "project",
    "site",
    "schedule_task",
    "start_time",
    "end_time",
    "status",
    "remarks",
}


def read_personal_schedules(input_path: str | Path | None = None, config_path: str | Path | None = None) -> dict[str, Any]:
    config = load_schedule_config(config_path)
    workbook_path = Path(input_path) if input_path else repo_root() / config["workbook"]["input_path"]
    workbook = load_workbook(workbook_path, data_only=True)
    detected_personal_sheets: list[str] = []
    records: list[dict[str, Any]] = []
    sheet_issues: list[dict[str, Any]] = []

    for sheet_name in workbook.sheetnames:
        if not is_personal_sheet(sheet_name):
            continue
        worksheet = workbook[sheet_name]
        header_map = {}
        for column_index in range(1, worksheet.max_column + 1):
            canonical = canonical_header(worksheet.cell(row=1, column=column_index).value)
            if canonical:
                header_map[canonical] = column_index
        detected_personal_sheets.append(sheet_name)
        missing_required_headers = sorted(REQUIRED_FIELDS - set(header_map))
        if missing_required_headers:
            sheet_issues.append(
                {
                    "severity": "ERROR",
                    "code": "MISSING_REQUIRED_HEADERS",
                    "message": f"Missing required columns: {', '.join(missing_required_headers)}.",
                    "person": normalise_text(sheet_name),
                    "source_sheet": sheet_name,
                    "source_row": 1,
                }
            )
            continue
        for row_index in range(2, worksheet.max_row + 1):
            raw_record = {field: None for field in KNOWN_FIELDS}
            for field, column_index in header_map.items():
                raw_record[field] = worksheet.cell(row=row_index, column=column_index).value
            if not any(normalise_text(value) for value in raw_record.values() if value is not None):
                continue
            raw_record.update(
                {
                    "person": normalise_text(sheet_name),
                    "source_sheet": sheet_name,
                    "source_row": row_index,
                    "source_workbook": workbook_path.name,
                }
            )
            records.append(raw_record)

    return {
        "records": records,
        "detected_personal_sheets": detected_personal_sheets,
        "sheet_issues": sheet_issues,
        "all_sheet_names": workbook.sheetnames,
        "input_workbook": str(workbook_path),
    }


def main() -> None:
    payload = read_personal_schedules()
    print(f"Detected personal sheets: {', '.join(payload['detected_personal_sheets'])}")
    print(f"Total records: {len(payload['records'])}")


if __name__ == "__main__":
    main()
