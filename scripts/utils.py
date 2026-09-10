from __future__ import annotations

import json
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SYSTEM_SHEETS = {
    "00_Master_Schedule",
    "01_Monthly_View",
    "02_Weekly_View",
    "03_Conflict_Report",
    "04_Availability",
    "Template",
    "Config",
}

VALID_STATUSES = {"Planned", "Confirmed", "Completed", "Cancelled", "Leave"}

HEADER_ALIASES = {
    "date": "date",
    "project": "project",
    "site": "site",
    "schedule / task": "schedule_task",
    "schedule/task": "schedule_task",
    "schedule": "schedule_task",
    "task": "schedule_task",
    "start time": "start_time",
    "end time": "end_time",
    "status": "status",
    "remarks": "remarks",
    "category": "category",
    "client": "client",
    "location": "location",
    "priority": "priority",
    "all day": "all_day",
    "remote / site / office": "work_mode",
    "confirmed / tentative": "confirmed_state",
    "last updated": "last_updated",
}

MASTER_COLUMNS = [
    "Date",
    "Day",
    "Person",
    "Project",
    "Site",
    "Schedule / Task",
    "Start Time",
    "End Time",
    "Status",
    "Category",
    "Priority",
    "Remarks",
    "All Day",
    "Source Sheet",
    "Source Row",
    "Last Updated",
]

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BANNER_FILL = PatternFill("solid", fgColor="FFF2CC")
BANNER_FONT = Font(bold=True, color="7F6000")
STATUS_FILLS = {
    "Planned": "D9EAF7",
    "Confirmed": "C6EFCE",
    "Completed": "E7E6E6",
    "Cancelled": "F4CCCC",
    "Leave": "FFE599",
    "CONFLICT": "F4CCCC",
    "AVAILABLE": "C6EFCE",
    "PARTIAL": "FFF2CC",
    "BOOKED": "F4CCCC",
    "LEAVE": "FFD966",
}


def repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def resolve_repo_path(path_value: str | None, default_relative: str) -> Path:
    relative = path_value or default_relative
    return repo_root() / relative


def load_schedule_config(config_path: str | Path | None = None) -> dict[str, Any]:
    path = Path(config_path) if config_path else repo_root() / "config" / "schedule_config.yaml"
    with path.open("r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle) or {}
    config.setdefault("workbook", {})
    config.setdefault("workday", {})
    config.setdefault("availability", {})
    config.setdefault("query", {})
    config["workbook"].setdefault("input_path", "Input/Personal_Schedule.xlsx")
    config["workbook"].setdefault("output_path", "Output/Master_Schedule.xlsx")
    config["workbook"].setdefault("normalized_json_path", "data/normalized_schedule.json")
    config["workbook"].setdefault("validation_report_path", "data/validation_report.json")
    config["workbook"].setdefault("summary_json_path", "data/build_summary.json")
    config["workday"].setdefault("start", "09:00")
    config["workday"].setdefault("end", "18:00")
    config["query"].setdefault("base_date", None)
    return config


def canonical_header(value: Any) -> str | None:
    if value is None:
        return None
    normalized = " ".join(str(value).strip().lower().split())
    return HEADER_ALIASES.get(normalized)


def normalise_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def cell_to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return bool(value)
    return normalise_text(value).lower() in {"true", "yes", "y", "1"}


def parse_excel_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalise_text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_excel_time(value: Any) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    text = normalise_text(value)
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def format_date(value: date | None) -> str:
    return value.isoformat() if value else ""


def format_time(value: time | None) -> str:
    return value.strftime("%H:%M") if value else ""


def is_personal_sheet(sheet_name: str) -> bool:
    return bool(sheet_name and sheet_name not in SYSTEM_SHEETS)


def build_display_time(record: dict[str, Any]) -> str:
    if record.get("all_day"):
        return "All Day"
    start_time = format_time(record.get("start_time"))
    end_time = format_time(record.get("end_time"))
    if start_time and end_time:
        return f"{start_time}-{end_time}"
    return start_time or end_time


def record_sort_key(record: dict[str, Any]) -> tuple[Any, ...]:
    start_time = record.get("start_time")
    start_sort = start_time if start_time else (time(0, 0) if record.get("all_day") else time(23, 59))
    return (record.get("date") or date.max, start_sort, record.get("person") or "")


def record_interval_minutes(record: dict[str, Any]) -> tuple[int, int] | None:
    if record.get("all_day"):
        return (0, 24 * 60)
    start_time = record.get("start_time")
    end_time = record.get("end_time")
    if not start_time or not end_time:
        return None
    start_minutes = start_time.hour * 60 + start_time.minute
    end_minutes = end_time.hour * 60 + end_time.minute
    return (start_minutes, end_minutes)


def overlap_minutes(first: tuple[int, int], second: tuple[int, int]) -> int:
    return max(0, min(first[1], second[1]) - max(first[0], second[0]))


def daterange(start_date: date, end_date: date):
    current = start_date
    while current <= end_date:
        yield current
        current += timedelta(days=1)


def parse_workday(config: dict[str, Any]) -> tuple[time, time]:
    start = parse_excel_time(config["workday"]["start"])
    end = parse_excel_time(config["workday"]["end"])
    if not start or not end:
        raise ValueError("Configured workday start/end must be valid HH:MM values.")
    return start, end


def workday_minutes(config: dict[str, Any]) -> tuple[int, int]:
    start, end = parse_workday(config)
    return (start.hour * 60 + start.minute, end.hour * 60 + end.minute)


def serialise_record(record: dict[str, Any]) -> dict[str, Any]:
    serialised = {}
    for key, value in record.items():
        if isinstance(value, date) and not isinstance(value, datetime):
            serialised[key] = value.isoformat()
        elif isinstance(value, time):
            serialised[key] = value.strftime("%H:%M")
        else:
            serialised[key] = value
    return serialised


def deserialise_record(record: dict[str, Any]) -> dict[str, Any]:
    restored = dict(record)
    restored["date"] = parse_excel_date(restored.get("date"))
    restored["start_time"] = parse_excel_time(restored.get("start_time"))
    restored["end_time"] = parse_excel_time(restored.get("end_time"))
    restored["all_day"] = cell_to_bool(restored.get("all_day"))
    last_updated = restored.get("last_updated")
    parsed_last_updated_date = parse_excel_date(last_updated)
    if parsed_last_updated_date and not isinstance(last_updated, str):
        restored["last_updated"] = parsed_last_updated_date.isoformat()
    return restored


def save_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)


def load_normalized_records(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    return [deserialise_record(record) for record in payload.get("records", [])]


def ensure_output_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def replace_sheet(workbook: Workbook, title: str, index: int | None = None):
    if title in workbook.sheetnames:
        workbook.remove(workbook[title])
    worksheet = workbook.create_sheet(title, index=index)
    return worksheet


def write_banner(worksheet, max_col: int):
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell = worksheet.cell(row=1, column=1, value="AUTO GENERATED / DO NOT EDIT MANUALLY")
    cell.fill = BANNER_FILL
    cell.font = BANNER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 24


def style_header_row(worksheet, row_index: int):
    for cell in worksheet[row_index]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def apply_table_style(worksheet, header_row: int, freeze_panes: str):
    worksheet.freeze_panes = freeze_panes
    if worksheet.max_row >= header_row:
        worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    for row in worksheet.iter_rows(min_row=header_row, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
        for cell in row:
            if cell.row != header_row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    for column_index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_index)
        max_length = 0
        for row_index in range(1, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 36)


def apply_date_time_formats(worksheet, header_row: int, date_headers: set[str] | None = None, time_headers: set[str] | None = None):
    date_headers = date_headers or {"Date"}
    time_headers = time_headers or {"Start Time", "End Time"}
    headers = {worksheet.cell(row=header_row, column=col).value: col for col in range(1, worksheet.max_column + 1)}
    for header, column in headers.items():
        if header in date_headers:
            for row in range(header_row + 1, worksheet.max_row + 1):
                worksheet.cell(row=row, column=column).number_format = "yyyy-mm-dd"
        if header in time_headers:
            for row in range(header_row + 1, worksheet.max_row + 1):
                worksheet.cell(row=row, column=column).number_format = "hh:mm"


def apply_status_formatting(worksheet, header_row: int, header_name: str, allowed_values: dict[str, str]):
    headers = {worksheet.cell(row=header_row, column=col).value: col for col in range(1, worksheet.max_column + 1)}
    if header_name not in headers or worksheet.max_row <= header_row:
        return
    column_letter = get_column_letter(headers[header_name])
    cell_range = f"{column_letter}{header_row + 1}:{column_letter}{worksheet.max_row}"
    for value, color in allowed_values.items():
        worksheet.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=PatternFill("solid", fgColor=color)),
        )


def summarise_record(record: dict[str, Any]) -> str:
    title = " / ".join(part for part in [record.get("project"), record.get("schedule_task")] if part and part != "UNKNOWN")
    if not title:
        title = record.get("schedule_task") or record.get("project") or "UNKNOWN"
    time_text = build_display_time(record)
    return "\n".join(part for part in [title, time_text] if part)
