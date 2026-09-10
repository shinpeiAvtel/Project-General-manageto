from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.build_master_schedule import build_output_artifacts
from scripts.calculate_availability import calculate_availability
from scripts.detect_conflicts import detect_conflicts
from scripts.read_personal_schedules import read_personal_schedules
from scripts.utils import load_schedule_config
from scripts.validate_schedule import validate_records

HEADERS = [
    "Date",
    "Project",
    "Site",
    "Schedule / Task",
    "Start Time",
    "End Time",
    "Status",
    "Remarks",
    "Category",
    "Priority",
    "All Day",
    "Last Updated",
]


def write_rows(worksheet, rows):
    worksheet.append(HEADERS)
    for row in rows:
        worksheet.append(row)


def create_sample_workbook(path: Path):
    workbook = Workbook()
    workbook.active.title = "00_Master_Schedule"
    for system_sheet in ["01_Monthly_View", "02_Weekly_View", "03_Conflict_Report", "04_Availability", "Template", "Config"]:
        workbook.create_sheet(system_sheet)
    shiraishi = workbook.create_sheet("Shiraishi")
    tanaka = workbook.create_sheet("Tanaka")
    suzuki = workbook.create_sheet("Suzuki")
    write_rows(
        shiraishi,
        [
            ["2026-09-10", "Project A", "Site A", "Kickoff", "09:00", "12:00", "Confirmed", "", "Field", "High", False, "2026-09-01"],
            ["2026-09-10", "Project B", "Site B", "Inspection", "10:00", "15:00", "Planned", "Conflict sample", "Field", "High", False, "2026-09-02"],
            ["2026-09-11", "Project A", "Site A", "Documentation", "13:00", "17:00", "Planned", "", "Office", "Medium", False, "2026-09-03"],
            ["2026-09-14", "", "Office", "Admin", "09:00", "18:00", "Confirmed", "", "Office", "Low", False, "2026-09-04"],
            ["2026-09-15", "", "", "Paid Leave", "", "", "Leave", "", "Leave", "High", True, "2026-09-05"],
        ],
    )
    write_rows(
        tanaka,
        [
            ["2026-09-10", "Project C", "Site East", "Survey", "09:00", "11:00", "Confirmed", "", "Field", "Medium", False, "2026-09-01"],
            ["2026-09-11", "Project C", "Site East", "Site Support", "09:00", "18:00", "Confirmed", "", "Field", "High", False, "2026-09-02"],
            ["2026-09-12", "Project D", "Site West", "Review", "10:00", "12:00", "Planned", "", "Office", "Medium", False, "2026-09-03"],
            ["2026-09-16", "Project A", "Site A", "Support", "14:00", "18:00", "Planned", "", "Field", "Medium", False, "2026-09-04"],
            ["2026-09-17", "", "Office", "Internal Meeting", "09:30", "11:00", "Confirmed", "", "Office", "Low", False, "2026-09-05"],
        ],
    )
    write_rows(
        suzuki,
        [
            ["2026-09-10", "Project E", "Remote", "Analysis", "08:30", "10:30", "Planned", "", "Remote", "Medium", False, "2026-09-01"],
            ["2026-09-12", "Project A", "Site A", "Installation", "09:00", "17:00", "Confirmed", "", "Field", "High", False, "2026-09-02"],
            ["2026-09-14", "Project A", "Site A", "Installation", "09:00", "12:00", "Confirmed", "", "Field", "High", False, "2026-09-03"],
            ["2026-09-14", "Project B", "Site B", "Follow-up", "14:00", "17:00", "Planned", "", "Field", "Medium", False, "2026-09-04"],
            ["2026-09-18", "", "", "Training", "", "", "Confirmed", "", "Training", "Low", True, "2026-09-05"],
        ],
    )
    workbook.save(path)


def test_read_personal_schedules_excludes_system_sheets(tmp_path):
    workbook_path = tmp_path / "Personal_Schedule.xlsx"
    create_sample_workbook(workbook_path)
    payload = read_personal_schedules(workbook_path)
    assert payload["detected_personal_sheets"] == ["Shiraishi", "Tanaka", "Suzuki"]
    assert len(payload["records"]) == 15
    assert payload["sheet_issues"] == []


def test_reader_reports_missing_required_headers(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Shiraishi"
    worksheet.append(["Date", "Project", "Site"])
    worksheet.append(["2026-09-10", "Project A", "Site A"])
    workbook_path = tmp_path / "missing_headers.xlsx"
    workbook.save(workbook_path)

    payload = read_personal_schedules(workbook_path)
    assert payload["records"] == []
    assert payload["sheet_issues"][0]["code"] == "MISSING_REQUIRED_HEADERS"


def test_validation_flags_invalid_rows_and_keeps_valid_rows():
    raw_records = [
        {"person": "Shiraishi", "source_sheet": "Shiraishi", "source_row": 2, "schedule_task": "Task", "date": "2026-09-10", "project": "P", "site": "S", "start_time": "09:00", "end_time": "08:00", "status": "Confirmed", "remarks": "", "category": "", "priority": "", "all_day": False, "last_updated": "2026-09-01", "source_workbook": "Personal_Schedule.xlsx"},
        {"person": "", "source_sheet": "Unknown", "source_row": 3, "schedule_task": "", "date": "bad-date", "project": "", "site": "", "start_time": "bad", "end_time": "10:00", "status": "Mystery", "remarks": "", "category": "", "priority": "", "all_day": False, "last_updated": "", "source_workbook": "Personal_Schedule.xlsx"},
        {"person": "Tanaka", "source_sheet": "Tanaka", "source_row": 4, "schedule_task": "Valid", "date": "2026-09-11", "project": "", "site": "", "start_time": "09:00", "end_time": "10:00", "status": "Confirmed", "remarks": "", "category": "", "priority": "", "all_day": False, "last_updated": "", "source_workbook": "Personal_Schedule.xlsx"},
    ]
    valid_records, issues = validate_records(raw_records)
    assert len(valid_records) == 1
    codes = {issue["code"] for issue in issues}
    assert {"TIME_ORDER", "MISSING_PERSON", "INVALID_DATE", "INVALID_START_TIME", "MISSING_TASK", "UNKNOWN_STATUS"}.issubset(codes)
    assert valid_records[0]["project"] == "UNKNOWN"
    assert valid_records[0]["site"] == "UNKNOWN"


def test_conflict_detection_and_availability():
    records = [
        {"person": "Shiraishi", "date": date(2026, 9, 10), "schedule_task": "A", "project": "P1", "start_time": None, "end_time": None, "all_day": True, "status": "Confirmed", "source_sheet": "Shiraishi", "source_row": 2},
        {"person": "Shiraishi", "date": date(2026, 9, 10), "schedule_task": "B", "project": "P2", "start_time": None, "end_time": None, "all_day": True, "status": "Confirmed", "source_sheet": "Shiraishi", "source_row": 3},
        {"person": "Tanaka", "date": date(2026, 9, 10), "schedule_task": "Leave", "project": "UNKNOWN", "start_time": None, "end_time": None, "all_day": True, "status": "Leave", "source_sheet": "Tanaka", "source_row": 2},
    ]
    conflicts = detect_conflicts(records)
    assert len(conflicts) == 1
    config = load_schedule_config()
    availability = calculate_availability(records, config)
    lookup = {(row["person"], row["date"]): row["availability"] for row in availability}
    assert lookup[("Shiraishi", date(2026, 9, 10))] == "BOOKED"
    assert lookup[("Tanaka", date(2026, 9, 10))] == "LEAVE"


def test_end_to_end_build_outputs_and_traceability(tmp_path, monkeypatch):
    repo_root = tmp_path
    for relative in ["Input", "Output", "data", "config"]:
        (repo_root / relative).mkdir(parents=True, exist_ok=True)
    create_sample_workbook(repo_root / "Input" / "Personal_Schedule.xlsx")
    (repo_root / "config" / "schedule_config.yaml").write_text(
        """
workbook:
  input_path: Input/Personal_Schedule.xlsx
  output_path: Output/Master_Schedule.xlsx
  normalized_json_path: data/normalized_schedule.json
  validation_report_path: data/validation_report.json
workday:
  start: \"09:00\"
  end: \"18:00\"
availability:
  start_date: \"2026-09-10\"
  end_date: \"2026-09-18\"
query:
  base_date: \"2026-09-10\"
""",
        encoding="utf-8",
    )
    monkeypatch.setattr("scripts.utils.repo_root", lambda: repo_root)
    monkeypatch.setattr("scripts.read_personal_schedules.repo_root", lambda: repo_root)
    monkeypatch.setattr("scripts.validate_schedule.repo_root", lambda: repo_root)
    monkeypatch.setattr("scripts.build_master_schedule.repo_root", lambda: repo_root)

    result = build_output_artifacts(repo_root / "config" / "schedule_config.yaml")
    assert len(result["records"]) == 15
    assert len(result["conflicts"]) == 1

    workbook = load_workbook(repo_root / "Output" / "Master_Schedule.xlsx")
    assert workbook.sheetnames == [
        "00_Master_Schedule",
        "01_Monthly_View",
        "02_Weekly_View",
        "03_Conflict_Report",
        "04_Availability",
    ]
    master = workbook["00_Master_Schedule"]
    headers = [master.cell(row=2, column=col).value for col in range(1, master.max_column + 1)]
    assert "Source Sheet" in headers
    assert "Source Row" in headers
    first_data_row = [master.cell(row=3, column=col).value for col in range(1, master.max_column + 1)]
    assert first_data_row[2] == "Suzuki"
    shiraishi_row = next(
        row
        for row in range(3, master.max_row + 1)
        if master.cell(row=row, column=3).value == "Shiraishi" and master.cell(row=row, column=6).value == "Kickoff"
    )
    assert master.cell(row=shiraishi_row, column=14).value == "Shiraishi"
    assert master.cell(row=shiraishi_row, column=15).value == 2
    monthly = workbook["01_Monthly_View"]
    assert monthly.cell(row=3, column=1).value == "Shiraishi"
    conflicts = workbook["03_Conflict_Report"]
    assert conflicts.max_row == 3
    assert conflicts.cell(row=3, column=1).value == "Shiraishi"
    availability = workbook["04_Availability"]
    assert any(availability.cell(row=row, column=3).value == "LEAVE" for row in range(3, availability.max_row + 1))


def test_end_to_end_invalid_config_dates_raise_clear_error(tmp_path, monkeypatch):
    repo_root = tmp_path
    for relative in ["Input", "Output", "data", "config"]:
        (repo_root / relative).mkdir(parents=True, exist_ok=True)
    create_sample_workbook(repo_root / "Input" / "Personal_Schedule.xlsx")
    (repo_root / "config" / "schedule_config.yaml").write_text(
        """
workbook:
  input_path: Input/Personal_Schedule.xlsx
  output_path: Output/Master_Schedule.xlsx
  normalized_json_path: data/normalized_schedule.json
  validation_report_path: data/validation_report.json
workday:
  start: \"09:00\"
  end: \"18:00\"
availability:
  start_date: \"2026/09/10\"
  end_date: \"not-a-date\"
query:
  base_date: \"2026-09-10\"
""",
        encoding="utf-8",
    )
    monkeypatch.setattr("scripts.utils.repo_root", lambda: repo_root)
    monkeypatch.setattr("scripts.read_personal_schedules.repo_root", lambda: repo_root)
    monkeypatch.setattr("scripts.validate_schedule.repo_root", lambda: repo_root)
    monkeypatch.setattr("scripts.build_master_schedule.repo_root", lambda: repo_root)

    try:
        build_output_artifacts(repo_root / "config" / "schedule_config.yaml")
        assert False, "Expected build_output_artifacts to fail on invalid config dates."
    except ValueError as exc:
        assert "availability.end_date" in str(exc)
