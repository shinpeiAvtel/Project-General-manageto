from datetime import date
from pathlib import Path

from openpyxl import Workbook

from scripts.calendar_schedule import _minutes, detect_conflicts, parse_calendar_sheet, build_calendar_sheet


def test_minutes_parser():
    assert _minutes('09:00-17:00') == (540, 1020)
    assert _minutes('ALL DAY') is None
    assert _minutes('bad') is None


def test_calendar_parse_and_conflict():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tester'
    build_calendar_sheet(ws, 'Tester', 2026, 9)
    # 2026-09-16 is in week block starting row 15, Wednesday = G:I, slots rows 16-18
    ws['G16'] = '10:00-16:00'
    ws['H16'] = 'Project A'
    ws['I16'] = 'T&C'
    ws['G17'] = '15:00-17:00'
    ws['H17'] = 'Project B'
    ws['I17'] = 'Design Review'
    records = parse_calendar_sheet(ws)
    assert len(records) == 2
    assert records[0]['person'] == 'Tester'
    assert records[0]['date'] == date(2026, 9, 16)
    conflicts = detect_conflicts(records)
    assert len(conflicts) == 1
    assert conflicts[0][0] == 'Tester'
