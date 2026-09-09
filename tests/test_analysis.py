from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from scripts.analyze_schedule import analyze_records
from scripts.generate_master_schedule import generate_master_schedule


def test_analysis_delay_conflict_stale_and_risk(tmp_path):
    records = [
        {
            "project_name": "NRT82",
            "project_manager": "Shiraishi",
            "task_id": "T1",
            "task_name": "Task1",
            "start_date": "2026-09-01",
            "end_date": "2026-09-05",
            "progress": 20,
            "milestone": True,
            "category": "Testing",
            "assigned_person": ["Tanaka"],
            "priority": "HIGH",
            "last_updated": "2026-08-20",
        },
        {
            "project_name": "HND10",
            "project_manager": "Tanaka",
            "task_id": "T2",
            "task_name": "Task2",
            "start_date": "2026-09-03",
            "end_date": "2026-09-07",
            "progress": 0,
            "milestone": False,
            "category": "T&C",
            "assigned_person": ["Tanaka"],
            "priority": "MEDIUM",
            "last_updated": "2026-08-25",
        },
    ]
    validation = [{"severity": "ERROR", "project_name": "NRT82"}]
    rules = {
        "milestone_categories": ["T&C", "Testing"],
        "risk_weights": {
            "delayed_task": 3,
            "high_priority_delayed_task": 5,
            "milestone_delay": 5,
            "resource_conflict": 2,
            "stale_warning": 2,
            "stale_high_warning": 4,
            "not_started_late": 2,
            "validation_error": 2,
        },
        "stale_update_threshold_days": {"warning": 7, "high_warning": 14},
        "risk_level_thresholds": {"low_max": 2, "medium_max": 6, "high_max": 10},
    }
    analysis = analyze_records(records, validation, rules, today=date(2026, 9, 10))

    assert any(t["delayed_status"] == "DELAYED" for t in analysis["delayed_tasks"])
    assert any(c["type"] == "POTENTIAL_CONFLICT" for c in analysis["resource_conflicts"])
    assert analysis["stale_projects"]["NRT82"] in {"WARNING", "HIGH WARNING"}
    assert analysis["risk_by_project"]["NRT82"]["score"] > 0
    assert len(analysis["milestone_candidates"]) == 1

    out_file = tmp_path / "Master_Schedule.xlsx"
    generate_master_schedule(analysis, validation, out_file)
    wb = load_workbook(out_file)
    assert "Master Schedule" in wb.sheetnames
    assert "Project Summary" in wb.sheetnames
    assert "Validation Errors" in wb.sheetnames
